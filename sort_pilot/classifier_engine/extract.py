from __future__ import annotations

import math
import mimetypes
import re
import struct
import threading
import time
import unicodedata
import zipfile
import zlib
from collections import Counter, OrderedDict
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from stop_words import get_stop_words

from .types import Feature, FeatureVector, path_id

IN_PROGRESS = {".crdownload", ".part", ".tmp", ".download"}
KO_PARTICLES = ("에서는", "으로", "에게", "에서", "부터", "까지", "처럼", "보다", "은", "는", "이", "가", "을", "를", "에", "의", "도", "와", "과")
STOP = frozenset(get_stop_words("en")) | frozenset(get_stop_words("ko")) | {"그리고", "합니다", "있는", "없는"}
MAX_BODY_TERMS = 160
MAX_COLLOCATIONS = 40
MAX_OCR_LAYOUT_ITEMS = 40
MIN_COLLOCATION_COUNT = 2
MIN_COLLOCATION_PMI = 2.0
COLLOCATION_PREFIXES = {2: "bi:", 3: "tri:"}
SEMANTIC_FEATURE_SOURCES = frozenset({"body", "ocr", "ocr_char", "obj", "pair"})
CONTENT_ANALYSIS_SUFFIXES = frozenset({
    ".txt", ".md", ".csv", ".rtf", ".pdf", ".hwp", ".hwpx", ".docx", ".odt", ".pptx", ".xlsx",
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic",
})
IMAGE_OCR_SUFFIXES = frozenset({".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"})
OCR_LOW_CONFIDENCE = 0.55
PDF_DIRECT_TEXT_MIN = 100
PDF_OCR_TEXT_MAX = 40
PDF_NORMAL_CHARACTER_RATIO = 0.85
PDF_IMAGE_AREA_RATIO = 0.70
PDF_OCR_DPI = 150
PDF_OCR_TIMEOUT_SECONDS = 3.0
PDF_SHORT_DOCUMENT_PAGES = 5
PDF_CACHE_ITEMS = 8
OCR_CHARACTER_HASH_BINS = 512
_KIWI = None
_OCR_LOCAL = threading.local()
_PDF_CACHE_LOCAL = threading.local()


@dataclass(frozen=True, slots=True)
class PDFExtraction:
    """Transient PDF text, OCR quality, metadata, and non-text numeric features."""

    text: str
    template_text: str
    layout_evidence: tuple[str, ...]
    metadata_features: tuple[Feature, ...]
    numeric_features: dict[str, float]
    quality: str
    ocr_confidence: float | None


@dataclass(frozen=True, slots=True)
class TransientOCRResult:
    """Small picklable OCR result that never includes page pixels or cropped images."""

    txts: tuple[str, ...]
    boxes: tuple[tuple[tuple[float, float], ...], ...] | None
    scores: tuple[float, ...]


class _PDFOCRWorker:
    """One reusable local OCR process with a hard per-page execution timeout."""

    def __init__(self) -> None:
        import multiprocessing

        context = multiprocessing.get_context("spawn")
        parent, child = context.Pipe(duplex=True)
        self._connection = parent
        self._process = context.Process(
            target=_pdf_ocr_worker_main,
            args=(child,),
            name="sort-pilot-pdf-ocr",
            daemon=True,
        )
        self._process.start()
        child.close()
        if not parent.poll(30.0):
            self.close()
            raise TimeoutError("Korean OCR worker initialization exceeded 30 seconds.")
        kind, detail = parent.recv()
        if kind != "ready":
            self.close()
            raise RuntimeError(f"Korean OCR worker initialization failed: {detail}")

    def recognize(self, image: object, timeout: float = 3.0) -> TransientOCRResult:
        """Return one compact result or terminate the worker at the hard timeout."""
        if not self._process.is_alive():
            raise RuntimeError("Korean OCR worker is not running.")
        self._connection.send(("ocr", image))
        if not self._connection.poll(timeout):
            self.close()
            raise TimeoutError(f"PDF OCR exceeded {timeout:.0f} seconds.")
        kind, value = self._connection.recv()
        if kind != "result" or not isinstance(value, TransientOCRResult):
            raise RuntimeError(f"PDF OCR worker failed: {value}")
        return value

    def close(self) -> None:
        """Stop this exact OCR worker and release its private pipe."""
        try:
            if self._process.is_alive():
                self._process.terminate()
            self._process.join(timeout=1.0)
        finally:
            self._connection.close()


def _pdf_ocr_worker_main(connection: object) -> None:
    """Initialize RapidOCR once, then handle transient page arrays until terminated."""
    try:
        engine = _ocr_engine()
        connection.send(("ready", None))
        while True:
            command, image = connection.recv()
            if command != "ocr":
                break
            try:
                connection.send(("result", _compact_ocr_result(engine(image))))
            except Exception as exc:
                connection.send(("error", type(exc).__name__))
    except EOFError:
        pass
    except Exception as exc:
        try:
            connection.send(("error", type(exc).__name__))
        except Exception:
            pass
    finally:
        connection.close()


def _compact_ocr_result(result: object) -> TransientOCRResult:
    """Drop images and retain only transient line geometry, text, and confidence."""
    raw_texts = getattr(result, "txts", None)
    texts = tuple(str(value) for value in (() if raw_texts is None else raw_texts))
    raw_scores = getattr(result, "scores", None)
    scores = tuple(float(value) for value in (() if raw_scores is None else raw_scores))
    raw_boxes = getattr(result, "boxes", None)
    boxes = None
    if raw_boxes is not None:
        boxes = tuple(
            tuple((float(point[0]), float(point[1])) for point in box)
            for box in raw_boxes
        )
    return TransientOCRResult(texts, boxes, scores)


def _pdf_ocr_worker() -> _PDFOCRWorker:
    """Reuse one PDF OCR worker for every call made from the current worker thread."""
    worker = getattr(_OCR_LOCAL, "pdf_worker", None)
    if worker is None or not worker._process.is_alive():
        worker = _PDFOCRWorker()
        _OCR_LOCAL.pdf_worker = worker
    return worker


def supports_content_analysis(path: Path) -> bool:
    """Return whether the bundled extractors can derive semantic content evidence."""
    return path.suffix.casefold() in CONTENT_ANALYSIS_SUFFIXES


def normalize_filename(path: Path) -> list[str]:
    """Normalize timestamp, copy, case, and separator patterns into tokens."""
    stem = re.sub(r"_?\d{8}_\d{6}", "", path.stem)
    stem = re.sub(r"\s*\(\d+\)$", "", stem)
    if stem.lower().rstrip("_") == "kakaotalk":
        return ["kakaotalk"]
    stem = re.sub(r"([a-z])([A-Z])", r"\1 \2", stem)
    return tokenize(re.sub(r"[_\-.]+", " ", stem))


def tokenize(text: str) -> list[str]:
    """Tokenize Korean and Latin text using Kiwi with a regex fallback."""
    global _KIWI
    try:
        if _KIWI is None:
            from kiwipiepy import Kiwi
            _KIWI = Kiwi()
        result = []
        for token in _KIWI.tokenize(text):
            if token.tag.startswith(("NN", "VV", "VA")) or token.tag in {"SL"}:
                form = token.form.lower()
                if form not in STOP and (len(form) >= 2 or re.search(r"[가-힣]", form)):
                    result.append(form)
        return result
    except ImportError:
        pass
    raw = re.findall(r"[가-힣]+|[A-Za-z][A-Za-z0-9]*", text.lower())
    result = []
    for token in raw:
        if re.search(r"[가-힣]", token):
            for particle in KO_PARTICLES:
                if token.endswith(particle) and len(token) > len(particle) + 1:
                    token = token[:-len(particle)]; break
        if token not in STOP and (len(token) >= 2 or re.search(r"[가-힣]", token)):
            result.append(token)
    return result


def collocations(tokens: list[str], n: int) -> Counter[str]:
    """Extract adjacent n-grams whose joint frequency exceeds chance (pointwise mutual information)."""
    if len(tokens) < n:
        return Counter()
    unigram_counts = Counter(tokens)
    total = len(tokens)
    ngram_counts = Counter(tuple(tokens[i:i + n]) for i in range(len(tokens) - n + 1))
    total_ngrams = sum(ngram_counts.values())
    kept: Counter[str] = Counter()
    for ngram, count in ngram_counts.items():
        if count < MIN_COLLOCATION_COUNT:
            continue
        joint = count / total_ngrams
        independent = 1.0
        for word in ngram:
            independent *= unigram_counts[word] / total
        pmi = math.log(joint / independent) if independent > 0 else 0.0
        if pmi >= MIN_COLLOCATION_PMI:
            kept[COLLOCATION_PREFIXES[n] + " ".join(ngram)] = count
    return kept


def _read_text(path: Path, limit: int) -> str:
    """Read bounded plain text using common Korean and Unicode encodings."""
    for encoding in ("utf-8-sig", "cp949", "utf-16"):
        try: return path.read_text(encoding=encoding)[:limit]
        except (UnicodeError, OSError): continue
    return ""


def _docx(path: Path, limit: int) -> str:
    """Extract bounded text directly from a DOCX document XML payload."""
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    return " ".join(node.text or "" for node in root.iter())[:limit]


def _odt(path: Path, limit: int) -> str:
    """Extract bounded paragraph and table text from an ODT content payload."""
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))
    return " ".join(node.text or "" for node in root.iter())[:limit]


def _archive(path: Path) -> str:
    """Extract representative top-level names from a ZIP archive."""
    if path.suffix.lower() != ".zip": return ""
    with zipfile.ZipFile(path) as archive:
        return " ".join(Path(name).parts[0] for name in archive.namelist()[:500])


def _pdf(path: Path, max_chars: int) -> PDFExtraction:
    """Return a transient cached result keyed only by path, size, and modification time."""
    stat = path.stat()
    key = (str(path.resolve()), stat.st_size, stat.st_mtime_ns)
    cache = getattr(_PDF_CACHE_LOCAL, "items", None)
    if cache is None:
        cache = OrderedDict()
        _PDF_CACHE_LOCAL.items = cache
    cached = cache.get(key)
    if cached is not None:
        cache.move_to_end(key)
        return cached
    result = _extract_pdf_pages(path, min(max_chars, 20_000))
    cache[key] = result
    cache.move_to_end(key)
    while len(cache) > PDF_CACHE_ITEMS:
        cache.popitem(last=False)
    return result


def _extract_pdf_pages(path: Path, max_chars: int) -> PDFExtraction:
    """Extract selected PDF pages, OCR scan-like pages, and retain numeric evidence."""
    import pymupdf

    page_texts: list[str] = []
    template_texts: list[str] = []
    layout_evidence: list[str] = []
    ocr_confidences: list[float] = []
    ocr_pages = 0
    failed_ocr_pages = 0
    scan_like_pages = 0
    image_ratios: list[float] = []
    extracted_characters = 0
    with pymupdf.open(path) as document:
        page_count = document.page_count
        selected = _pdf_page_indices(page_count)
        for page_index in selected:
            page = document.load_page(page_index)
            embedded = page.get_text("text") or ""
            normal_ratio = _normal_character_ratio(embedded)
            image_ratio = _pdf_image_ratio(page)
            image_ratios.append(image_ratio)
            direct_text = (
                len(embedded.strip()) >= PDF_DIRECT_TEXT_MIN
                and normal_ratio >= PDF_NORMAL_CHARACTER_RATIO
                and image_ratio < PDF_IMAGE_AREA_RATIO
            )
            should_ocr = not direct_text and (
                len(embedded.strip()) < PDF_OCR_TEXT_MAX
                or image_ratio >= PDF_IMAGE_AREA_RATIO
                or normal_ratio < PDF_NORMAL_CHARACTER_RATIO
            )
            scan_like_pages += int(should_ocr)
            page_ocr_text = ""
            page_template_text = ""
            if should_ocr and ocr_pages == 0:
                ocr_pages += 1
                try:
                    result, confidence = _ocr_pdf_page(page, PDF_OCR_DPI)
                except (TimeoutError, RuntimeError, OSError):
                    failed_ocr_pages += 1
                else:
                    page_ocr_text, page_layout = _layout_aware_ocr(result)
                    page_template_text = _original_ocr_text(result)
                    layout_evidence.extend(page_layout)
                    if confidence is not None:
                        ocr_confidences.append(confidence)
                    if not page_ocr_text or confidence is None or confidence < OCR_LOW_CONFIDENCE:
                        failed_ocr_pages += 1
            combined = _deduplicate_text((embedded, page_ocr_text), max_chars)
            combined_template = _deduplicate_text((embedded, page_template_text), max_chars)
            page_texts.append(combined)
            template_texts.append(combined_template)
            extracted_characters += len(combined)
        metadata = document.metadata or {}
        metadata_features = _pdf_metadata_features(metadata)

    selected_count = len(selected)
    quality = (
        "failed"
        if selected_count == 0 or (ocr_pages and failed_ocr_pages == ocr_pages and not extracted_characters)
        else "low_confidence"
        if failed_ocr_pages
        else "ok"
    )
    confidence = min(ocr_confidences) if ocr_confidences else None
    return PDFExtraction(
        text=_deduplicate_text(page_texts, max_chars),
        template_text=_deduplicate_text(template_texts, max_chars),
        layout_evidence=tuple(dict.fromkeys(layout_evidence))[:MAX_OCR_LAYOUT_ITEMS],
        metadata_features=metadata_features,
        numeric_features={
            "pdf_page_count": float(page_count),
            "pdf_sampled_page_count": float(selected_count),
            "pdf_scan_ratio": (
                0.0 if not selected_count else scan_like_pages / selected_count
            ),
            "pdf_image_ratio": 0.0 if not image_ratios else sum(image_ratios) / len(image_ratios),
            "pdf_text_characters_per_page": (
                0.0 if not selected_count else extracted_characters / selected_count
            ),
        },
        quality=quality,
        ocr_confidence=confidence,
    )


def _pdf_page_indices(page_count: int) -> tuple[int, ...]:
    """Inspect every short-PDF page or first three, middle, and last for long PDFs."""
    if page_count <= 0:
        return ()
    if page_count <= PDF_SHORT_DOCUMENT_PAGES:
        return tuple(range(page_count))
    return tuple(dict.fromkeys((0, 1, 2, page_count // 2, page_count - 1)))


def _normal_character_ratio(text: str) -> float:
    """Measure printable letters, numbers, punctuation, symbols, and whitespace."""
    if not text:
        return 0.0
    normal = sum(
        unicodedata.category(character)[0] in {"L", "M", "N", "P", "S", "Z"}
        or character in "\r\n\t"
        for character in text
    )
    return normal / len(text)


def _pdf_image_ratio(page: object) -> float:
    """Return a bounded sum of image bounding-box area over the page area."""
    page_area = max(0.0, float(page.rect.width) * float(page.rect.height))
    if not page_area:
        return 0.0
    area = 0.0
    try:
        for value in page.get_image_info(xrefs=True):
            bbox = value.get("bbox")
            if bbox is None or len(bbox) != 4:
                continue
            left, top, right, bottom = map(float, bbox)
            area += max(0.0, right - left) * max(0.0, bottom - top)
    except (AttributeError, TypeError, ValueError):
        return 0.0
    return min(1.0, area / page_area)


def _ocr_pdf_page(page: object, dpi: int = PDF_OCR_DPI) -> tuple[object, float | None]:
    """Render one page and hard-stop its local OCR worker after three seconds."""
    pixmap = page.get_pixmap(dpi=dpi, alpha=False)
    import numpy as np

    image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
        pixmap.height,
        pixmap.width,
        pixmap.n,
    )
    if pixmap.n >= 3:
        image = image[:, :, :3][:, :, ::-1].copy()
    result = _pdf_ocr_worker().recognize(image, timeout=PDF_OCR_TIMEOUT_SECONDS)
    return result, _ocr_confidence(result)


def _pdf_metadata_features(metadata: dict) -> tuple[Feature, ...]:
    """Represent title, author, and keywords only as weak lexical features."""
    features: list[Feature] = []
    for field in ("title", "author", "keywords"):
        for token in tokenize(str(metadata.get(field) or ""))[:20]:
            features.append(Feature(f"{field}:{token}", "pdf_meta"))
    return tuple(features)


def _deduplicate_text(values: object, limit: int) -> str:
    """Merge bounded text blocks while dropping normalized duplicates and substrings."""
    kept: list[str] = []
    normalized: list[str] = []
    for raw in values:
        value = re.sub(r"\s+", " ", str(raw or "")).strip()
        key = unicodedata.normalize("NFKC", value).casefold()
        if not key or any(key == prior or key in prior for prior in normalized):
            continue
        kept = [text for text, prior in zip(kept, normalized, strict=True) if prior not in key]
        normalized = [prior for prior in normalized if prior not in key]
        kept.append(value)
        normalized.append(key)
        if sum(len(text) + 1 for text in kept) >= limit:
            break
    return "\n".join(kept)[:limit]


def _hwp(path: Path, limit: int) -> str:
    """Extract bounded text from local HWP 5.0 or HWPX documents."""
    return _hwpx(path, limit) if path.suffix.casefold() == ".hwpx" else _hwp5(path, limit)


def _hwpx(path: Path, limit: int) -> str:
    """Read paragraph text from the ZIP/XML HWPX representation."""
    with zipfile.ZipFile(path) as archive:
        section_names = sorted(
            (
                name
                for name in archive.namelist()
                if re.fullmatch(r"(?i)Contents/section\d+\.xml", name)
            ),
            key=lambda name: int(re.search(r"\d+", Path(name).stem).group()),
        )
        if not section_names:
            raise ValueError("HWPX document has no section XML.")
        values: list[str] = []
        length = 0
        for name in section_names:
            root = ElementTree.fromstring(archive.read(name))
            for element in root.iter():
                if element.tag.rsplit("}", 1)[-1] != "t" or not element.text:
                    continue
                text = element.text.strip()
                if not text:
                    continue
                values.append(text)
                length += len(text) + 1
                if length >= limit:
                    return "\n".join(values)[:limit]
        return "\n".join(values)[:limit]


def _hwp5(path: Path, limit: int) -> str:
    """Parse compressed HWP 5.0 PARA_TEXT records through its OLE streams."""
    import olefile

    if not olefile.isOleFile(str(path)):
        raise ValueError("HWP 5.0 document is not an OLE compound file.")
    with olefile.OleFileIO(str(path)) as document:
        if not document.exists("FileHeader"):
            raise ValueError("HWP 5.0 document has no FileHeader stream.")
        header = document.openstream("FileHeader").read()
        if len(header) < 40 or not header.startswith(b"HWP Document File"):
            raise ValueError("HWP 5.0 header is invalid.")
        flags = struct.unpack_from("<I", header, 36)[0]
        if flags & 0x02:
            raise ValueError("Password-protected HWP text cannot be extracted.")
        if flags & 0x04:
            raise ValueError("Distribution-protected HWP text cannot be extracted.")
        section_names = sorted(
            (
                "/".join(parts)
                for parts in document.listdir()
                if len(parts) == 2
                and parts[0] == "BodyText"
                and re.fullmatch(r"Section\d+", parts[1])
            ),
            key=lambda name: int(name.removeprefix("BodyText/Section")),
        )
        values: list[str] = []
        length = 0
        for name in section_names:
            data = document.openstream(name).read()
            if flags & 0x01:
                data = zlib.decompress(data, -15)
            for tag_id, payload in _hwp_records(data):
                if tag_id != 67:
                    continue
                text = _decode_hwp_para_text(payload).strip()
                if text:
                    values.append(text)
                    length += len(text) + 1
                    if length >= limit:
                        return "\n".join(values)[:limit]
        if values:
            return "\n".join(values)[:limit]
        if document.exists("PrvText"):
            return document.openstream("PrvText").read().decode("utf-16le")[:limit]
        return ""


def _hwp_records(data: bytes):
    """Yield bounded `(tag_id, payload)` pairs including extended-size records."""
    offset = 0
    while offset + 4 <= len(data):
        header = struct.unpack_from("<I", data, offset)[0]
        offset += 4
        tag_id = header & 0x3FF
        size = (header >> 20) & 0xFFF
        if size == 0xFFF:
            if offset + 4 > len(data):
                break
            size = struct.unpack_from("<I", data, offset)[0]
            offset += 4
        end = offset + size
        if end > len(data):
            break
        yield tag_id, data[offset:end]
        offset = end


def _decode_hwp_para_text(data: bytes) -> str:
    """Drop fixed-width HWP controls while retaining displayed Unicode text."""
    values: list[str] = []
    ordinary = bytearray()

    def flush() -> None:
        if ordinary:
            values.append(ordinary.decode("utf-16le", errors="replace"))
            ordinary.clear()

    offset = 0
    extended_controls = {
        1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19,
        20, 21, 22, 23,
    }
    while offset + 2 <= len(data):
        code = struct.unpack_from("<H", data, offset)[0]
        if code >= 32:
            ordinary.extend(data[offset : offset + 2])
            offset += 2
            continue
        flush()
        if code in {10, 13}:
            values.append("\n")
        elif code in {9, 24, 30, 31}:
            values.append(" ")
        offset += 16 if code in extended_controls else 2
    flush()
    return re.sub(r"[ \t]+", " ", "".join(values)).replace(" \n", "\n")


def _pptx(path: Path, limit: int) -> str:
    """Extract bounded text from PowerPoint slide text frames."""
    from pptx import Presentation
    deck = Presentation(path)
    return " ".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text_frame"))[:limit]


def _xlsx(path: Path, limit: int) -> str:
    """Extract bounded sheet names and cell values from a workbook."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    values = []
    try:
        for sheet in book.worksheets:
            values.append(sheet.title)
            for row in sheet.iter_rows(max_row=50, max_col=20, values_only=True):
                values.extend(str(value) for value in row if value is not None)
                if sum(map(len, values)) >= limit: break
    finally: book.close()
    return " ".join(values)[:limit]


def _image_features(path: Path) -> tuple[str, list[Feature]]:
    """Route an image and derive camera, aspect, and color metadata."""
    from PIL import Image, ImageStat
    with Image.open(path) as image:
        exif = image.getexif(); width, height = image.size
        has_camera = bool(exif.get(271) or exif.get(272))
        known = {(1920, 1080), (2560, 1440), (1366, 768), (3840, 2160)}
        route = "photo" if has_camera else "screenshot" if image.size in known or path.suffix.lower() == ".png" else "ambiguous"
        features = [Feature(f"image:{route}", "meta"), Feature("aspect:" + ("portrait" if height > width else "square" if abs(width-height)/max(width,height)<.1 else "landscape"), "meta")]
        if has_camera: features.append(Feature("exif:camera_present", "meta"))
        sample = image.convert("RGB"); sample.thumbnail((80, 80)); mean = ImageStat.Stat(sample).mean
        features.append(Feature("color:warm" if mean[0] > mean[2] * 1.08 else "color:cool", "meta"))
        return route, features


def _ocr_engine():
    """Return one verified local Korean RapidOCR engine per worker thread."""
    from rapidocr import RapidOCR
    from .ocr import KoreanOCRAssets

    engine = getattr(_OCR_LOCAL, "engine", None)
    if engine is None:
        assets = KoreanOCRAssets.load()
        engine = RapidOCR(params=assets.rapidocr_params())
        _OCR_LOCAL.engine = engine
    return engine


def _ocr_image_input(path: Path):
    """Apply EXIF orientation and normalize supported images before local OCR."""
    if not path.is_file():
        return str(path)
    from PIL import Image, ImageOps
    import numpy as np

    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source)
        if getattr(image, "n_frames", 1) > 1:
            image.seek(0)
        rgb = image.convert("RGB")
        return np.asarray(rgb, dtype=np.uint8)[:, :, ::-1].copy()


def _original_ocr_text(result: object) -> str:
    """Return bounded detector order for template/layout semantics only."""
    raw_values = getattr(result, "txts", None)
    values = () if raw_values is None else raw_values
    return " ".join(
        str(value).strip()
        for value in values
        if str(value).strip()
    )[:20_000]


def _ocr_confidence(result: object) -> float | None:
    """Return mean finite RapidOCR recognition confidence without persisting text."""
    values: list[float] = []
    raw_scores = getattr(result, "scores", None)
    for raw in (() if raw_scores is None else raw_scores):
        try:
            value = float(raw)
        except (TypeError, ValueError):
            continue
        if math.isfinite(value) and 0.0 <= value <= 1.0:
            values.append(value)
    return None if not values else sum(values) / len(values)


def _hashed_character_features(text: str) -> list[Feature]:
    """Create privacy-preserving OCR typo features for character TF-IDF downstream."""
    import hashlib

    normalized = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text).casefold()).strip()
    counts: Counter[int] = Counter()
    for size in (3, 4, 5):
        for offset in range(max(0, len(normalized) - size + 1)):
            value = normalized[offset : offset + size]
            if not value.strip():
                continue
            digest = hashlib.blake2b(value.encode("utf-8"), digest_size=8).digest()
            counts[int.from_bytes(digest, "big") % OCR_CHARACTER_HASH_BINS] += 1
    return [
        Feature(f"char:{index:03d}", "ocr_char", float(count))
        for index, count in counts.most_common(OCR_CHARACTER_HASH_BINS)
    ]


def _layout_aware_ocr(result: object) -> tuple[str, tuple[str, ...]]:
    """Order OCR lines by page columns and retain bounded transient layout evidence."""
    raw_texts = getattr(result, "txts", None)
    texts = tuple(str(value).strip() for value in (() if raw_texts is None else raw_texts))
    texts = tuple(value for value in texts if value)
    boxes = getattr(result, "boxes", None)
    if not texts:
        return "", ()
    if boxes is None or len(boxes) != len(texts):
        ordered = texts
    else:
        lines: list[tuple[str, float, float, float, float]] = []
        try:
            for text, box in zip(texts, boxes, strict=True):
                xs = tuple(float(point[0]) for point in box)
                ys = tuple(float(point[1]) for point in box)
                if not xs or not ys or not all(math.isfinite(value) for value in (*xs, *ys)):
                    raise ValueError
                lines.append((text, min(xs), min(ys), max(xs), max(ys)))
        except (TypeError, ValueError, IndexError):
            ordered = texts
        else:
            page_left = min(line[1] for line in lines)
            page_right = max(line[3] for line in lines)
            midpoint = (page_left + page_right) / 2.0
            tolerance = max(1.0, (page_right - page_left) * 0.04)
            left = [line for line in lines if line[3] < midpoint + tolerance]
            right = [
                line
                for line in lines
                if line not in left and line[1] > midpoint - tolerance
            ]
            spanning = [line for line in lines if line not in left and line not in right]
            if len(left) >= 2 and len(right) >= 2:
                first_column_y = min(line[2] for line in (*left, *right))
                headers = [line for line in spanning if line[2] <= first_column_y]
                footers = [line for line in spanning if line not in headers]
                ordered_lines = (
                    sorted(headers, key=lambda line: (line[2], line[1]))
                    + sorted(left, key=lambda line: (line[2], line[1]))
                    + sorted(right, key=lambda line: (line[2], line[1]))
                    + sorted(footers, key=lambda line: (line[2], line[1]))
                )
            else:
                ordered_lines = sorted(lines, key=lambda line: (line[2], line[1]))
            ordered = tuple(line[0] for line in ordered_lines)
    natural_text = " ".join(ordered)[:20_000]
    evidence: list[str] = []
    for line in ordered:
        for value in (line[:120], re.sub(r"\s+", "", line)[:120]):
            if value and value not in evidence:
                evidence.append(value)
            if len(evidence) >= MAX_OCR_LAYOUT_ITEMS:
                break
        if len(evidence) >= MAX_OCR_LAYOUT_ITEMS:
            break
    return natural_text, tuple(evidence)


def _ocr_evidence(path: Path) -> tuple[str, list[Feature], tuple[str, ...]]:
    """Extract bounded OCR natural text, lexical features, and transient layout evidence."""
    result = _ocr_engine()(_ocr_image_input(path))
    original_text = _original_ocr_text(result)
    natural_text, layout_evidence = _layout_aware_ocr(result)
    _OCR_LOCAL.last_text = natural_text
    _OCR_LOCAL.last_template_text = original_text
    _OCR_LOCAL.last_layout_evidence = layout_evidence
    _OCR_LOCAL.last_confidence = _ocr_confidence(result)
    return (
        natural_text,
        [
            *[Feature(token, "ocr") for token in tokenize(natural_text)[:MAX_BODY_TERMS]],
            *_hashed_character_features(natural_text),
        ],
        layout_evidence,
    )


def _ocr(path: Path) -> list[Feature]:
    """Return OCR lexical features for the legacy extractor interface."""
    return _ocr_evidence(path)[1]


def extract(path: Path, max_content_mb=200, max_chars=20_000) -> FeatureVector:
    """Build a bounded local feature vector from filename, content, and media."""
    started = time.perf_counter(); stat = path.stat(); features = []
    features.extend(Feature(t, "filename") for t in normalize_filename(path))
    features.append(Feature(path.suffix.lower().lstrip(".") or "no_ext", "ext"))
    if re.match(r"(?i)^KakaoTalk_\d{8}_\d{6}", path.name): features.append(Feature("kakao_export", "meta"))
    partial = stat.st_size > max_content_mb * 1024 * 1024
    text = ""
    template_text = ""
    ocr_layout_evidence: tuple[str, ...] = ()
    numeric_features: dict[str, float] = {}
    extraction_quality = "ok"
    ocr_confidence: float | None = None
    suffix = path.suffix.lower()
    if not partial:
        try:
            if suffix in {".txt", ".md", ".csv", ".rtf"}: text = _read_text(path, max_chars)
            elif suffix == ".pdf":
                pdf = _pdf(path, max_chars)
                text = pdf.text
                template_text = pdf.template_text
                ocr_layout_evidence = pdf.layout_evidence
                numeric_features = pdf.numeric_features
                extraction_quality = pdf.quality
                ocr_confidence = pdf.ocr_confidence
                features.extend(pdf.metadata_features)
            elif suffix in {".hwp", ".hwpx"}: text = _hwp(path, max_chars)
            elif suffix == ".docx": text = _docx(path, max_chars)
            elif suffix == ".odt": text = _odt(path, max_chars)
            elif suffix == ".pptx": text = _pptx(path, max_chars)
            elif suffix == ".xlsx": text = _xlsx(path, max_chars)
            elif suffix == ".zip": text = _archive(path)
        except Exception:
            partial = True
            extraction_quality = "failed"
    if not template_text:
        template_text = text
    body_tokens = tokenize(text)
    for token, count in Counter(body_tokens).most_common(MAX_BODY_TERMS):
        features.append(Feature(token, "body", float(count)))
    for n in (2, 3):
        for ngram, count in collocations(body_tokens, n).most_common(MAX_COLLOCATIONS):
            features.append(Feature(ngram, "body", float(count)))
    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    route = "text" if text else "metadata"
    is_image = mime.startswith("image/") or suffix in IMAGE_OCR_SUFFIXES
    if is_image:
        try:
            route, image_features = _image_features(path); features.extend(image_features)
            if suffix in IMAGE_OCR_SUFFIXES:
                _OCR_LOCAL.last_text = ""
                _OCR_LOCAL.last_template_text = ""
                _OCR_LOCAL.last_layout_evidence = ()
                _OCR_LOCAL.last_confidence = None
                features.extend(_ocr(path))
                ocr_text = getattr(_OCR_LOCAL, "last_text", "")
                template_ocr_text = getattr(_OCR_LOCAL, "last_template_text", "")
                ocr_layout_evidence = getattr(_OCR_LOCAL, "last_layout_evidence", ())
                ocr_confidence = getattr(_OCR_LOCAL, "last_confidence", None)
                if not text:
                    text = ocr_text[:max_chars]
                    template_text = template_ocr_text[:max_chars]
                if not ocr_text:
                    extraction_quality = "failed"
                elif ocr_confidence is not None and ocr_confidence < OCR_LOW_CONFIDENCE:
                    extraction_quality = "low_confidence"
        except Exception:
            partial = True
            route = "image"
            extraction_quality = "failed"
    elif suffix == ".pdf" and numeric_features.get("pdf_scan_ratio", 0.0) > 0.0:
        features.extend(_hashed_character_features(text))
    if (
        supports_content_analysis(path)
        and not is_image
        and suffix != ".pdf"
        and not text.strip()
    ):
        extraction_quality = "failed"
    return FeatureVector(
        path_id(path),
        str(path),
        stat.st_size,
        features,
        partial,
        route,
        {"total": (time.perf_counter()-started)*1000},
        natural_text=text,
        template_natural_text=template_text,
        ocr_layout_evidence=ocr_layout_evidence,
        numeric_features=numeric_features,
        extraction_quality=extraction_quality,
        ocr_confidence=ocr_confidence,
    )


def is_processable(path: Path, exclusions=()) -> bool:
    """Reject directories, partial downloads, office locks, and excluded paths."""
    import fnmatch
    return path.is_file() and path.suffix.lower() not in IN_PROGRESS and not path.name.startswith("~$") and not any(fnmatch.fnmatch(str(path), p) for p in exclusions)


def is_stable(path: Path, interval=2.0) -> bool:
    """Check that file size and modification time remain unchanged."""
    try:
        first = path.stat(); time.sleep(interval); second = path.stat()
        if (first.st_size, first.st_mtime_ns) != (second.st_size, second.st_mtime_ns): return False
        with path.open("rb"): pass
        return True
    except (OSError, PermissionError): return False
